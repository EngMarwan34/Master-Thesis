"""برنامج تعبئة النماذج: تعبّي الفراغات في المتصفح ← يطلع النموذج جاهز (PDF + Excel).
التشغيل:  python app.py   ثم افتح  http://localhost:8000
المتطلبات: pip install openpyxl   (و LibreOffice لإخراج PDF، وإلا يُنتج Excel فقط)"""
import datetime as dt, json, os, shutil, subprocess, sys, tempfile, uuid
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import quote, unquote
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "output")
os.makedirs(OUT, exist_ok=True)
DAYS = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
MARK = "✔"


def d(s):
    try:
        return dt.date.fromisoformat(s) if s else None
    except ValueError:
        return None


def fmt(x):
    return x.strftime("%Y/%m/%d") if x else ""


def num(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def hours(a, b):
    try:
        t1 = dt.datetime.strptime(a, "%H:%M"); t2 = dt.datetime.strptime(b, "%H:%M")
        h = (t2 - t1).total_seconds() / 3600
        return round(h + 24 if h < 0 else h, 2)
    except (TypeError, ValueError):
        return None


def fill_overtime(ws, f):
    ws["A10"] = f"عن الفتره من :- {fmt(d(f['from']))} الي {fmt(d(f['to']))}"
    ws["A12"] = f"الاداره :- {f['branch']}" if f["branch"] else None
    ws["A13"] = f"القسم :- {f['dept']}"
    ws["A14"] = f"اسم المشروع :- {f['project']}"
    ws["D14"] = f"رقم المشروع :- {f['project_no']}"
    for i, r in enumerate(f["rows"][:14]):
        row = 17 + i
        dd = d(r.get("date"))
        ws[f"B{row}"] = r.get("name") or None
        ws[f"C{row}"] = DAYS[dd.weekday()] if dd else None
        ws[f"D{row}"] = fmt(dd) or None
        ws[f"E{row}"] = r.get("start") or None
        ws[f"F{row}"] = r.get("end") or None
        h = num(r.get("hours")) if r.get("hours") else hours(r.get("start"), r.get("end"))
        ws[f"G{row}"] = h
    ws["G31"] = "=SUM(G17:G30)"


def fill_meals(ws, f):
    ws["A13"] = f"اسم القسم :-  {f['dept']}          التاريخ :-  {fmt(d(f['date']))}"
    ws["A14"] = f"اسم العميل: {f['client']}"
    ws["A15"] = f"اسم المشروع:  {f['project']}"
    ws["A16"] = f"رقم المشروع : {f['project_no']}"
    for i, r in enumerate(f["rows"][:10]):
        row = 20 + i
        a, b = d(r.get("from")), d(r.get("to"))
        days = num(r.get("days")) or ((b - a).days + 1 if a and b else None)
        ws[f"C{row}"] = r.get("name") or None
        ws[f"D{row}"] = a; ws[f"E{row}"] = b
        for c in "DE":
            ws[f"{c}{row}"].number_format = "yyyy/mm/dd"
        ws[f"F{row}"] = days
        ws[f"G{row}"] = num(r.get("price")) if r.get("price") else num(f.get("price"))
        ws[f"H{row}"] = f"=G{row}*F{row}" if days and ws[f"G{row}"].value is not None else None
    ws["H30"] = "=SUM(H20:H29)"


def fill_purchase(ws, f):
    dd = d(f["date"])
    ws["D14"] = DAYS[dd.weekday()] if dd else None
    ws["D15"] = fmt(dd) or None
    ws["D16"] = f["dept"]; ws["D17"] = f["supplier"]
    for base, c in ((19, f["c1"]), (24, f["c2"])):
        ws[f"D{base}"] = c["project"] or None
        ws[f"D{base+1}"] = c["quote"] or None
        ws[f"D{base+2}"] = c["sys"] or None
        r = base + 3
        ws[f"D{r}"] = num(c["advance"]) if c["advance"] else None
        ws[f"F{r}"] = MARK if c["type"] == "credit" else None
        ws[f"H{r}"] = MARK if c["type"] == "cash" else None
    total = 0
    for i, r in enumerate(f["rows"][:5]):
        row = 31 + i
        q, p = num(r.get("qty")), num(r.get("price"))
        ws[f"C{row}"] = r.get("desc") or None
        ws[f"G{row}"] = q; ws[f"H{row}"] = p
        ws[f"I{row}"] = f"=G{row}*H{row}" if q is not None and p is not None else None
    ws["I36"] = "=SUM(I31:I35)"


FORMS = {"overtime": ("overtime.xlsx", fill_overtime, "كشف_الساعات_الاضافية"),
         "meals": ("meals.xlsx", fill_meals, "بدل_وجبة"),
         "purchase": ("purchase.xlsx", fill_purchase, "طلب_شراء")}


def to_pdf(xlsx):
    exe = shutil.which("soffice") or shutil.which("libreoffice") or next(
        (p for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                     r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                     "/Applications/LibreOffice.app/Contents/MacOS/soffice") if os.path.exists(p)), None)
    if not exe:
        return None
    # مجلد ملف شخصي مؤقت حتى لا يتعارض مع LibreOffice المفتوح
    prof = tempfile.mkdtemp()
    try:
        subprocess.run([exe, f"-env:UserInstallation=file://{prof}".replace("file://C:", "file:///C:"),
                        "--headless", "--convert-to", "pdf", "--outdir", OUT, xlsx],
                       check=True, capture_output=True, timeout=120)
    except Exception:
        return None
    finally:
        shutil.rmtree(prof, ignore_errors=True)
    pdf = os.path.splitext(xlsx)[0] + ".pdf"
    return pdf if os.path.exists(pdf) else None


def generate(kind, data):
    tpl, fn, label = FORMS[kind]
    wb = openpyxl.load_workbook(os.path.join(BASE, "templates", tpl))
    fn(wb.worksheets[0], data)
    name = f"{label}_{dt.datetime.now():%Y%m%d_%H%M%S}"
    path = os.path.join(OUT, name + ".xlsx")
    wb.save(path)
    pdf = to_pdf(path)
    return name, bool(pdf)


class H(BaseHTTPRequestHandler):
    def log_message(self, *a): pass

    def send(self, code, body, ctype="text/html; charset=utf-8", extra=None):
        if isinstance(body, str): body = body.encode()
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        for k, v in (extra or {}).items(): self.send_header(k, v)
        self.end_headers(); self.wfile.write(body)

    def do_GET(self):
        p = unquote(self.path.split("?")[0])
        if p == "/":
            with open(os.path.join(BASE, "index.html"), encoding="utf-8") as fh:
                return self.send(200, fh.read())
        if p.startswith("/files/"):
            fn = os.path.basename(p)
            fp = os.path.join(OUT, fn)
            if os.path.exists(fp):
                ct = "application/pdf" if fn.endswith(".pdf") else \
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                with open(fp, "rb") as fh:
                    return self.send(200, fh.read(), ct,
                                     {"Content-Disposition": f"inline; filename*=UTF-8''{quote(fn)}"})
        self.send(404, "not found")

    def do_POST(self):
        try:
            n = int(self.headers.get("Content-Length", 0))
            body = json.loads(self.rfile.read(n))
            kind = self.path.strip("/")
            name, has_pdf = generate(kind, body)
            self.send(200, json.dumps({"name": name, "pdf": has_pdf}), "application/json")
        except Exception as e:
            self.send(500, json.dumps({"error": str(e)}), "application/json")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    srv = ThreadingHTTPServer(("127.0.0.1", port), H)
    print(f"افتح المتصفح على http://localhost:{port}  (إيقاف: Ctrl+C)")
    if "--no-browser" not in sys.argv:
        webbrowser.open(f"http://localhost:{port}")
    srv.serve_forever()
